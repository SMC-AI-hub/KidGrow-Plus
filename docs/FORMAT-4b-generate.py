"""Regenerate FORMAT-4b-security-verification-report.docx from the repo record.

Usage:  python3 docs/FORMAT-4b-generate.py
Output: docs/FORMAT-4b-security-verification-report.docx
"""
from docx import Document
from docx.shared import Pt

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10.5)

def h1(t):
    doc.add_heading(t, level=1)

def h2(t):
    doc.add_heading(t, level=2)

def para(t, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(t)
    r.bold = bold
    return p

def bullets(items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")

def table(headers, rows):
    tb = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tb.style = "Light Grid Accent 1"
    for j, htxt in enumerate(headers):
        tb.rows[0].cells[j].text = htxt
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            tb.rows[i + 1].cells[j].text = val
    doc.add_paragraph()
    return tb

# ---------- Title ----------
doc.add_heading("KidGrow-Plus (NeoKid Care) — Security Verification Report", level=0)
para("Item 4b · Formatted sign-off deliverable · Generated 2026-10-23 from the repo record "
     "(regenerate any time with docs/FORMAT-4b-generate.py).")
table(["Field", "Value"], [
    ["Branch / commits", "staging, 02679df (Phase 2) … 944c73a (dashboard 8/8)"],
    ["Environments", "Staging bubnidgejmbzkevxvthl · Prod jmlvcrekttbzibqitftj"],
    ["Migrations live", "0007 + 0008 on BOTH envs · 0009 on staging only (prod gated to 2026-11-05)"],
    ["Suites", "isolation 146/146 · cloud-offline 598/598 · full npm test exit 0"],
])

# ---------- 1. Phase 1 findings ----------
h1("1. Phase 1 findings (audit 2026-10-13, re-verified through 2026-10-23)")
para("Repo result stands: 16 RLS-enabled tables, 0 views, 21 locked DEFINER functions, "
     "0 always-true policies, 0 npm vulnerabilities, no committed secrets (148-commit scan).")
table(["ID", "Finding", "Severity", "Status"], [
    ["S1", "No tenant isolation in legacy backend", "Critical", "Done — 0001 tenant core + isolation suite green"],
    ["S2", "Staff PINs plaintext, synced to cloud", "High", "Done — salted SHA-256 + lockout (36/36 tests)"],
    ["S3", "Supabase JWT in localStorage (XSS steals session)", "High", "Open — mitigations in place, residual documented"],
    ["S4", "Paywall UI-only (devtools bypass)", "Medium", "Accepted — server enforcement via M3 lifecycle"],
    ["S5", "Service worker caches all same-origin GET", "Medium", "Done — /api/* bypassed + asserted"],
    ["S6", "No security headers / CSP", "Medium", "Done — HSTS/CSP asserted, enforced-header proof"],
    ["S7", "Encrypted-backup 4-char passwords", "Medium", "Done — minimum 8 + asserted"],
    ["S8", "No MFA / lockout / session revoke", "Medium", "Partial — TOTP + lockout done; 0009 adds server aal2 (staging)"],
    ["S9", "API functions lack schema validation + rate limits", "Low", "Planned — before any PHI API"],
    ["S10", "Webhook redelivery duplicates rows", "Low", "Planned — dedupe key (money-effect idempotent)"],
    ["S11", "No package.json/lockfile (supply-chain drift)", "Low", "Done — lockfile + npm test wired"],
])
para("Verified non-issues: no eval/document.write; OSV zero vulns for vendored "
     "supabase-js 2.117.3 / qrcode-generator 1.4.4; consistent output encoding in report paths.", bold=False)

# ---------- 2. Phase 2 ----------
h1("2. Phase 2 verification (0007 role governance + 0008 TRUNCATE revoke)")
h2("2.1 Migration 0007 — staging VERIFIED 2026-10-21/22")
bullets([
    "Apply: Success. Census at verify: 2 tenants / 4 memberships / 4 patients.",
    "Fingerprint 2/7/2 (functions / block_delete triggers / CHECKs).",
    "E-block E1–E5 all PASS (readonly deny x2, owner DELETE 0 + row intact, owner positive controls x2).",
    "Q1=0 (no DELETE policies) · Q2=0 (no PUBLIC/anon EXECUTE) · Q5 = 0007 set exactly · Q8 = 7/7 triggers.",
    "Q7 baseline recorded: full grants incl. TRUNCATE/TRIGGER to anon+authenticated on all 16 tables — RLS is the live gate.",
])
h2("2.2 Migration 0007 — prod VERIFIED 2026-10-22")
bullets([
    "Pre-flight green (census 0/0/0, fingerprint 0/0/0, 3 auth users). Apply: Success.",
    "Fingerprint 2/7/2, E1–E5 PASS, census+Q1+Q2 clean, Q5 exact.",
    "Prod Q7: uniform 5-privilege pattern, never had TRUNCATE/TRIGGER (staging-only drift).",
])
h2("2.3 Migration 0008 — VERIFIED both envs 2026-10-22")
bullets([
    "Staging: TRUNCATE/TRIGGER grants 64 -> 0 rows; E-block 5/5 PASS (no behavior change).",
    "Prod: pre-check already 0; apply no-op Success; post 0; E-block 5/5 PASS.",
    "Suite 134/134 at the time. Latent TRUNCATE-bypasses-RLS finding RESOLVED.",
])

# ---------- 3. P3.1 ----------
h1("3. P3.1 staging verification (0009 aal2 enforcement + grace UX)")
bullets([
    "Commit 7cdc84f. Staging apply: Success (2026-10-22).",
    "0009 gates: owner/admin/doctor WRITES (clinical via tenant_can_write + privileged-actor policies) "
    "and privileged audit READS (audit_events, break_glass_events) on aal2.",
    "Deliberately ungated: nurse writes, auditor reads, audit/break-glass INSERTs, tenant mint, "
    "profiles/subscription reads, membership reads, accept_invite/ensure_profile.",
    "App grace UX (ships first): MFA_GRACE_UNTIL 2026-11-05; amber optional-verify banner before, "
    "red required banner after; 42501 failures route to verification, not raw errors.",
    "Suites: isolation 146/146 (12 new 0009 asserts) · cloud-offline 598/598 (3 grace asserts) · full run exit 0.",
    "PROD 0009 GATED until on/after 2026-11-05 (warn-then-enforce; locked decision).",
])
table(["E-check", "Expectation", "Staging result"], [
    ["E1", "readonly direct INSERT -> 42501", "PASS"],
    ["E2", "readonly sync-RPC -> 42501", "PASS"],
    ["E3", "owner DELETE -> 0 rows, intact", "PASS"],
    ["E4", "owner aal2 INSERT -> allowed", "PASS"],
    ["E5", "owner aal2 sync-RPC -> inserted", "PASS"],
    ["E6", "owner aal1 INSERT -> 42501 (P3.1)", "PASS"],
    ["E7", "audit reads: aal2 sees 1, aal1 sees 0", "PASS"],
])
table(["Step B baseline (corrected 2026-10-22)", "Value"], [
    ["my_aal2 grantees", "postgres (owner) + authenticated + service_role — 3 rows, none PUBLIC/anon"],
])

# ---------- 4. Locked decisions ----------
h1("4. Locked Phase 3 decisions")
table(["Key", "Decision"], [
    ["(a) aal2 roles", "owner + admin + doctor"],
    ["(b) Scope", "writes + sensitive reads"],
    ["(c) Rollout", "14-day warn-then-enforce (grace until 2026-11-05)"],
    ["(d/e) Tenant minting", "P3.2 owns t_ins; cap 1 tenant/user"],
    ["(f) Confirm-email", "ON (staging enabled 2026-10-23; prod already on)"],
    ["(g) L4 framing", "X-Frame-Options: DENY"],
    ["(h) M3 PHI encryption", "full design inside Phase 3"],
])

# ---------- 5. Dashboard values ----------
h1("5. Dashboard values (copied from Supabase dashboard 2026-10-23)")
table(["#", "Value", "Staging", "Prod"], [
    ["1", "Project ref", "bubnidgejmbzkevxvthl", "jmlvcrekttbzibqitftj"],
    ["2", "Region", "ap-south-1", "(to confirm)"],
    ["3", "Postgres", "17.11.0.003", "(to confirm)"],
    ["4", "MFA/TOTP", "Enabled", "(to confirm)"],
    ["5", "Confirm-email", "ON (enabled 2026-10-23)", "ON (already on)"],
    ["6", "Backup plan", "None — Free plan", "Likely none — FREE org (CONFIRM; go-live blocker)"],
    ["7", "Earliest restore", "n/a", "(to confirm)"],
    ["8", "JWT expiry", "3600 (default kept)", "3600 (assumed default)"],
])
bullets([
    "Both projects migrated HS256 -> ECC P-256 signing keys ~Oct 15 (upgrade). "
    "Previous keys revokable now — recommended on both projects.",
    "Prod Data API posture: 16/16 tables exposed, 0/23 functions, max rows 1000; "
    "'Automatically expose new tables' ON — recommended OFF.",
])

# ---------- 6. Clinical sign-offs ----------
h1("6. Clinical sign-off status (F1 / F9 / F25)")
bullets([
    "F1 dose sheet (30 rows) verified CURRENT against js/drugs.js 30/30 — ready to sign, UNSIGNED.",
    "F9 growth spot-check sheet (21 cells, WHO + INTERGROWTH) created — UNSIGNED.",
    "F25 NICU case-log sheet (10 real de-identified cases) created — UNSIGNED.",
    "Evidence locker evidence/m7/ scaffolded; filing + retention (>=3y) rules recorded.",
    "Rule: no real prescriptions/patients until all three are clinician-signed and filed.",
])

# ---------- 7. Open items ----------
h1("7. Open items and go-live blockers")
table(["#", "Item", "Owner", "State"], [
    ["1", "Prod paid backup/PITR plan", "Owner", "BLOCKER — confirm + upgrade before go-live"],
    ["2", "Prod 0009 apply + verify", "Agent+Owner", "Gated to on/after 2026-11-05"],
    ["3", "F1/F9/F25 clinician sign-offs", "Clinician", "Sheets ready; awaiting validator"],
    ["4", "PAT rotation (token traveled via chat)", "Owner", "2 min in GitHub settings"],
    ["5", "Custom SMTP on prod", "Owner", "Before go-live (invites/resets deliverability)"],
    ["6", "Revoke previous HS256 JWT keys", "Owner", "One click per project"],
    ["7", "Auto-expose new tables -> OFF", "Owner", "One toggle per project"],
    ["8", "M7 gates (pentest/UAT/restore/DPDP/go-no-go)", "Humans", "After sign-offs"],
    ["9", "anon/service_role -> new API-key format", "Agent", "Later Phase 3 item"],
])

# ---------- 8. Evidence ----------
h1("8. Evidence index (repo paths)")
bullets([
    "docs/security/phase2-completion.md — Phase 2 + P3.1 completion record",
    "docs/security/phase2-live-verify.sql — Part 1 inventory + Part 2 E-block v2 (E1–E7)",
    "docs/security/phase3-plan.md — Phase 3 draft plan with decision checklist",
    "docs/security/dashboard-values.md — dashboard values 8/8 + posture notes",
    "supabase/migrations/0007_role_governance.sql, 0008_revoke_truncate_trigger.sql, 0009_aal_enforcement.sql",
    "tests/tenant-isolation.test.js (146), tests/cloud-offline.test.js (598)",
    "docs/F1-DOSE-SIGNOFF.md, docs/F9-GROWTH-SPOTCHECK.md, docs/F25-NICU-CASELOG.md, evidence/m7/",
    "evidence/ui/ — dashboard + settings screenshots",
])

# ---------- 9. Sign-offs ----------
h1("9. Sign-off")
para("By signing, reviewers confirm the results above were witnessed (live output pasted back) "
     "and the open items in section 7 are understood.")
table(["Role", "Name", "Signature", "Date"], [
    ["Security reviewer (agent-assisted)", "", "", ""],
    ["Clinic owner", "", "", ""],
    ["Clinical validator (F1/F9/F25)", "", "", ""],
])

doc.save("docs/FORMAT-4b-security-verification-report.docx")
print("WROTE docs/FORMAT-4b-security-verification-report.docx")
