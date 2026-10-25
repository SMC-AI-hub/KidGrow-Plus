#!/usr/bin/env python3
"""Strict validator: app WHO_REF vs (1) official WHO expanded tables, (2) CDC WHOref_d daily LMS."""
import json, csv, math
import openpyxl

REPO = '/home/user/kidgrow-pro'
src = open(f'{REPO}/js/ref-who.js').read()
app = json.loads(src.split('const WHO_REF = ', 1)[1].rsplit(';', 1)[0])

def lmsVal(M, S, L, z):
    if L == 0:
        return M * math.exp(S * z)
    b = 1 + L * S * z
    return M * (b ** (1 / L)) if b > 0 else float('nan')

MAP = [('wfa', 'wfa'), ('lhfa', 'lhfa'), ('bfa', 'bmifa'), ('hcfa', 'hcfa'),
       ('acfa', 'muacfa'), ('wfl', 'wfl'), ('wfh', 'wfh')]
SEX = [('b', '_b'), ('g', '_g')]
ZCOLS = {-3: 'SD3neg', -2: 'SD2neg', -1: 'SD1neg', 0: 'SD0', 1: 'SD1', 2: 'SD2', 3: 'SD3'}

print('=== SOURCE 1: official WHO expanded tables (cdn.who.int) vs app ===')
grand = []
for off, apl in MAP:
    for sx, suf in SEX:
        wb = openpyxl.load_workbook(f'/tmp/whox/{off}{suf}.xlsx', read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
        hdr = [str(c) for c in rows[0]]
        zi = {z: hdr.index(c) for z, c in ZCOLS.items()}
        T = app[apl][sx]
        xs = T.get('x') or T.get('age')
        lut = {x: i for i, x in enumerate(xs)}
        n = 0
        md = {'L': (0, None), 'M': (0, None), 'S': (0, None), 'SD': (0, None)}
        for r in rows[1:]:
            x = r[0]
            if x is None or not isinstance(x, (int, float)):
                continue
            if apl in ('wfl', 'wfh'):
                if abs(x * 2 - round(x * 2)) > 1e-9:
                    continue
                x = round(x, 1)
            else:
                if abs(x - round(x)) > 1e-9:
                    continue
                x = int(round(x))
            if x not in lut:
                continue
            i = lut[x]
            L, M, S = T['L'][i], T['M'][i], T['S'][i]
            oL, oM, oS = r[1], r[2], r[3]
            n += 1
            for k, a, o in (('L', L, oL), ('M', M, oM), ('S', S, oS)):
                d = abs(a - o)
                if d > md[k][0]:
                    md[k] = (d, x)
            for z, col in zi.items():
                ov = r[col]
                if ov is None:
                    continue
                av = lmsVal(M, S, L, z)
                d = abs(av - ov)
                if d > md['SD'][0]:
                    md['SD'] = (d, (x, z))
        flag = 'FLAG' if (md['L'][0] > 2e-4 or md['M'][0] > 2e-4 or md['S'][0] > 2e-4 or md['SD'][0] > 0.002) else 'ok'
        print(f'{apl}/{sx}: n={n} max|dL|={md["L"][0]:.2e}@{md["L"][1]} max|dM|={md["M"][0]:.2e}@{md["M"][1]} '
              f'max|dS|={md["S"][0]:.2e}@{md["S"][1]} max|dSD|={md["SD"][0]:.2e}@{md["SD"][1]} [{flag}]')
        grand.append((f'{apl}/{sx}', flag, md))

print()
print('=== SOURCE 2: CDC WHOref_d.csv (daily LMS, 0-2y) vs app ===')
cdc = list(csv.DictReader(open('/tmp/WHOref_d.csv')))
print('cdc rows:', len(cdc), '| cols:', [c for c in cdc[0].keys() if c.strip()][:30])
CMAP = [('_wei', 'wfa'), ('_len', 'lhfa'), ('_bmi', 'bmifa'), ('_headc', 'hcfa'), ('_armc', 'muacfa')]
for prefix, apl in CMAP:
    for sex, sx in (('1', 'b'), ('2', 'g')):
        T = app[apl][sx]
        xs = T.get('x') or T.get('age')
        lut = {x: i for i, x in enumerate(xs)}
        n = 0
        md = {'L': (0, None), 'M': (0, None), 'S': (0, None)}
        for r in cdc:
            if r['sex'] != sex or (r.get('_denom') or '').strip() not in ('', 'forage'):
                continue
            dstr = (r['_agedays'] or '').strip()
            if not dstr:
                continue
            x = int(float(dstr))
            if x not in lut:
                continue
            try:
                oL = float(r[prefix + '_l']); oM = float(r[prefix + '_m']); oS = float(r[prefix + '_s'])
            except (ValueError, TypeError, KeyError):
                continue
            i = lut[x]
            n += 1
            for k, a, o in (('L', T['L'][i], oL), ('M', T['M'][i], oM), ('S', T['S'][i], oS)):
                dd = abs(a - o)
                if dd > md[k][0]:
                    md[k] = (dd, x)
        flag = 'FLAG' if (md['L'][0] > 2e-4 or md['M'][0] > 2e-4 or md['S'][0] > 2e-4) else 'ok'
        print(f'{apl}/{sx}: n={n} max|dL|={md["L"][0]:.2e}@{md["L"][1]} max|dM|={md["M"][0]:.2e}@{md["M"][1]} max|dS|={md["S"][0]:.2e}@{md["S"][1]} [{flag}]')
